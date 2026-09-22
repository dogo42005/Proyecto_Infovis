"""
Consolida los datasets crudos de SUBTEL y CONASET (data/cruda/) en tablas
limpias por año (data/procesada/).

Uso:
    python procesar_datos.py
"""
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
CRUDA = BASE / "cruda"
PROCESADA = BASE / "procesada"
PROCESADA.mkdir(exist_ok=True)


def normalizar(texto):
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(texto).strip().lower())
    return texto.encode("ascii", "ignore").decode()


def leer_por_etiqueta(nombre_archivo, predicado, columnas, max_col=9):
    """Recorre las hojas 2000..2025 (una por año) de un archivo CONASET y
    extrae, en cada una, la primera fila cuya etiqueta (columna A) cumple
    `predicado`. `columnas` mapea nombre_final -> índice de columna."""
    wb = load_workbook(CRUDA / nombre_archivo, read_only=True, data_only=True)
    registros = []
    for nombre_hoja in wb.sheetnames:
        if not nombre_hoja.strip().isdigit():
            continue
        anio = int(nombre_hoja)
        ws = wb[nombre_hoja]
        for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_col, values_only=True):
            if predicado(normalizar(fila[0])):
                registro = {"anio": anio}
                registro.update({nombre: fila[idx] for nombre, idx in columnas.items()})
                registros.append(registro)
                break
    wb.close()
    return pd.DataFrame(registros).sort_values("anio").reset_index(drop=True)


# ---------------------------------------------------------------------------
# SUBTEL — abonados móviles (serie principal de "uso del celular")
# ---------------------------------------------------------------------------
def procesar_abonados_moviles():
    wb = load_workbook(CRUDA / "1_ABONADOS_MOVILES-JUN26-170826.xlsx", read_only=True, data_only=True)
    ws = wb["3.1.Abonados"]
    registros = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6, values_only=True):
        if isinstance(fila[1], int) and fila[2] == "Dic":
            registros.append({
                "anio": fila[1],
                "abonados_moviles": fila[3],
                "crecimiento_anual": fila[4],
                "penetracion_cada_100_hab": fila[5],
            })
    wb.close()
    df = pd.DataFrame(registros).sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "subtel_abonados_moviles.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# SUBTEL — líneas fijas y tráfico móvil (secundarios, no entran al merge)
# ---------------------------------------------------------------------------
def procesar_lineas_fijas():
    wb = load_workbook(CRUDA / "1_SERIES_LINEAS_TELEFONICAS-JUN26-170826.xlsx", read_only=True, data_only=True)
    ws = wb["1.1_Mensual_Nac"]
    registros = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6, values_only=True):
        if isinstance(fila[1], int) and fila[2] == "Dic":
            registros.append({
                "anio": fila[1],
                "lineas_fijas": fila[3],
                "penetracion_cada_100_hab": fila[5],
            })
    wb.close()
    df = pd.DataFrame(registros).sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "subtel_lineas_fijas.csv", index=False)
    return df


def procesar_trafico_movil():
    wb = load_workbook(CRUDA / "2_SERIES_TRAFICO_MOVILES-JUN26-170826.xlsx", read_only=True, data_only=True)
    ws = wb["4.1. Total_Móvil"]
    registros = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5, values_only=True):
        if isinstance(fila[1], int):
            registros.append({
                "anio": fila[1],
                "trafico_voz_miles_minutos": fila[3],
                "trafico_voz_miles_llamadas": fila[4],
            })
    wb.close()
    df = pd.DataFrame(registros).sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "subtel_trafico_voz_movil.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# CONASET — atropellos (serie principal de "consecuencia")
# ---------------------------------------------------------------------------
def procesar_atropello():
    columnas = {
        "siniestros": 2,
        "fallecidos": 3,
        "lesionados_graves": 4,
        "lesionados_menos_graves": 5,
        "lesionados_leves": 6,
        "total_lesionados": 7,
    }
    df = leer_por_etiqueta(
        "Tiposiniestro2000-2025.xlsx",
        lambda t: t == "total atropello",
        columnas,
    )
    df.to_csv(PROCESADA / "conaset_atropello.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# CONASET — víctimas peatones por calidad de participante (cruce/respaldo)
# ---------------------------------------------------------------------------
def procesar_peatones_calidad():
    columnas = {
        "fallecidos": 1,
        "lesionados_graves": 2,
        "lesionados_menos_graves": 3,
        "lesionados_leves": 4,
        "total_lesionados": 5,
    }
    df = leer_por_etiqueta(
        "Calidaddeparticipantes2000-2025.xlsx",
        lambda t: t == "peatones",
        columnas,
    )
    df.to_csv(PROCESADA / "conaset_peatones_calidad.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# CONASET — causas atribuibles al peatón (imprudencia + alcohol)
# La etiqueta "Imprudencia de/del Peatón" cambia de redacción entre años,
# por eso el predicado matchea por substring normalizado, no texto exacto.
# ---------------------------------------------------------------------------
def procesar_causas_peaton():
    columnas = {
        "siniestros": 2,
        "fallecidos": 3,
        "lesionados_graves": 4,
        "lesionados_menos_graves": 5,
        "lesionados_leves": 6,
        "total_lesionados": 7,
    }
    imprudencia = leer_por_etiqueta(
        "Causas_desgregadas_conaset_carabineros2000-2025.xlsx",
        lambda t: t.startswith("total imprudencia") and "peat" in t,
        columnas,
    ).add_prefix("imprudencia_peaton_").rename(columns={"imprudencia_peaton_anio": "anio"})

    alcohol = leer_por_etiqueta(
        "Causas_desgregadas_conaset_carabineros2000-2025.xlsx",
        lambda t: t.startswith("total alcohol") and "peat" in t,
        columnas,
    ).add_prefix("alcohol_peaton_").rename(columns={"alcohol_peaton_anio": "anio"})

    df = imprudencia.merge(alcohol, on="anio", how="outer").sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "conaset_causas_peaton.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# CONASET — zona de ocurrencia (urbana / rural)
# ---------------------------------------------------------------------------
def procesar_zona_ocurrencia():
    columnas = {
        "siniestros": 1,
        "fallecidos": 2,
        "lesionados_graves": 3,
        "lesionados_menos_graves": 4,
        "lesionados_leves": 5,
        "total_lesionados": 6,
    }
    urbana = leer_por_etiqueta(
        "Zonadeocurrencia2000-2025.xlsx", lambda t: t == "urbana", columnas
    ).add_prefix("urbana_").rename(columns={"urbana_anio": "anio"})

    rural = leer_por_etiqueta(
        "Zonadeocurrencia2000-2025.xlsx", lambda t: t == "rural", columnas
    ).add_prefix("rural_").rename(columns={"rural_anio": "anio"})

    df = urbana.merge(rural, on="anio", how="outer").sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "conaset_zona_ocurrencia.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# CONASET — evolución general 1972-2025 (tabla única, contexto de largo plazo)
# ---------------------------------------------------------------------------
def procesar_evolucion_general():
    wb = load_workbook(CRUDA / "EvoluciónsiniestrostransitoChile-1972-2025.xlsx", read_only=True, data_only=True)
    ws = wb["Chile (1972-2025)"]
    registros = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10, values_only=True):
        if isinstance(fila[0], int):
            registros.append({
                "anio": fila[0],
                "siniestros": fila[1],
                "fallecidos": fila[2],
                "lesionados_graves": fila[3],
                "lesionados_menos_graves": fila[4],
                "lesionados_leves": fila[5],
                "total_lesionados": fila[6],
                "total_victimas": fila[7],
                "tasa_motorizacion": fila[8],
                "vehiculos_cada_100_hab": fila[9],
            })
    wb.close()
    df = pd.DataFrame(registros).sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "conaset_evolucion_general.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# CONASET — tasa de fallecidos cada 10.000 vehículos (tabla única, secundaria)
# ---------------------------------------------------------------------------
def procesar_tasa_fallecidos_vehiculos():
    wb = load_workbook(CRUDA / "Tasa-de-fallecidos-cada-10.000-vehículos-1990-2025.xlsx", read_only=True, data_only=True)
    ws = wb["Tasa"]
    registros = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4, values_only=True):
        if isinstance(fila[0], int):
            registros.append({
                "anio": fila[0],
                "fallecidos": fila[1],
                "parque_vehicular": fila[2],
                "tasa_fallecidos_cada_10000_vehiculos": fila[3],
            })
    wb.close()
    df = pd.DataFrame(registros).sort_values("anio").reset_index(drop=True)
    df.to_csv(PROCESADA / "conaset_tasa_fallecidos_vehiculos.csv", index=False)
    return df


# ---------------------------------------------------------------------------
# Tabla principal para la visualización: une celular (SUBTEL) y atropellos
# (CONASET) por año, 2000-2025.
# ---------------------------------------------------------------------------
def construir_dataset_principal(abonados, atropello, peatones_calidad):
    df = (
        abonados[["anio", "abonados_moviles", "penetracion_cada_100_hab"]]
        .merge(
            atropello[["anio", "siniestros", "fallecidos", "total_lesionados"]].rename(columns={
                "siniestros": "atropello_siniestros",
                "fallecidos": "atropello_fallecidos",
                "total_lesionados": "atropello_total_lesionados",
            }),
            on="anio", how="inner",
        )
        .merge(
            peatones_calidad[["anio", "fallecidos", "total_lesionados"]].rename(columns={
                "fallecidos": "peatones_fallecidos_calidad",
                "total_lesionados": "peatones_total_lesionados_calidad",
            }),
            on="anio", how="inner",
        )
        .sort_values("anio")
        .reset_index(drop=True)
    )
    df.to_csv(PROCESADA / "dataset_principal.csv", index=False)
    return df


def main():
    abonados = procesar_abonados_moviles()
    procesar_lineas_fijas()
    procesar_trafico_movil()
    atropello = procesar_atropello()
    peatones_calidad = procesar_peatones_calidad()
    procesar_causas_peaton()
    procesar_zona_ocurrencia()
    procesar_evolucion_general()
    procesar_tasa_fallecidos_vehiculos()
    principal = construir_dataset_principal(abonados, atropello, peatones_calidad)

    print(f"abonados_moviles:   {len(abonados)} filas ({abonados['anio'].min()}-{abonados['anio'].max()})")
    print(f"atropello:          {len(atropello)} filas ({atropello['anio'].min()}-{atropello['anio'].max()})")
    print(f"dataset_principal:  {len(principal)} filas ({principal['anio'].min()}-{principal['anio'].max()})")
    print(f"\nArchivos escritos en: {PROCESADA}")


if __name__ == "__main__":
    main()
